"""GRN pages (GRN with Invoice, GRN post IQC, Manage Lots "GRN & Lot" tab):

  * Date of Invoice From/To filter is applied server-side
    (?date_from=YYYY-MM-DD&date_to=YYYY-MM-DD) to the table and to the export.
    GRNImport.invoice_date is free text, so several formats must parse.
  * Row checkboxes + "Export GRN": POST /grn/export (xlsx) carries every GRN
    and Lot field, honours the selection (ids), and without a selection exports
    every GRN matching the current filters.
"""
import io
import pathlib
import subprocess
import sys
import uuid
from types import SimpleNamespace

from tests.test_iqc_new_user import _login, make_user  # noqa: F401  (fixture)

ROOT = str(pathlib.Path(__file__).resolve().parent.parent)


# ── pure helpers ────────────────────────────────────────────────────────────

def test_parse_invoice_date_handles_free_text_formats():
    from datetime import date
    from routers.grn import parse_invoice_date
    assert parse_invoice_date("2026-03-05") == date(2026, 3, 5)
    assert parse_invoice_date("05-03-2026") == date(2026, 3, 5)
    assert parse_invoice_date("05/03/2026") == date(2026, 3, 5)
    assert parse_invoice_date("05 Mar 2026") == date(2026, 3, 5)
    assert parse_invoice_date("Mar 5, 2026") == date(2026, 3, 5)
    assert parse_invoice_date("") is None
    assert parse_invoice_date(None) is None
    assert parse_invoice_date("not a date") is None


def _g(num, inv_date):
    return SimpleNamespace(grn_number=num, invoice_date=inv_date)


def test_filter_by_invoice_date_range_inclusive_and_open_ended():
    from routers.grn import filter_grns_by_invoice_date
    rows = [_g("A", "2026-01-10"), _g("B", "15/02/2026"), _g("C", "2026-03-20"),
            _g("D", None), _g("E", "garbage")]
    nums = lambda r: [g.grn_number for g in r]  # noqa: E731
    # no / invalid bounds -> untouched (undated rows stay visible)
    assert nums(filter_grns_by_invoice_date(rows, "", "")) == list("ABCDE")
    assert nums(filter_grns_by_invoice_date(rows, "bad", "")) == list("ABCDE")
    # both bounds inclusive; free-text date parsed; undated/unparseable excluded
    assert nums(filter_grns_by_invoice_date(rows, "2026-01-10", "2026-02-15")) == ["A", "B"]
    # open-ended
    assert nums(filter_grns_by_invoice_date(rows, "2026-02-01", "")) == ["B", "C"]
    assert nums(filter_grns_by_invoice_date(rows, "", "2026-01-31")) == ["A"]
    assert nums(filter_grns_by_invoice_date(rows, "2027-01-01", "2027-12-31")) == []


def test_export_columns_cover_every_grn_and_lot_field():
    from models.grn_import import GRNImport
    from models.lot import Lot
    from routers.grn import GRN_EXPORT_COLUMNS
    headers = " | ".join(h for h, _ in GRN_EXPORT_COLUMNS).lower()
    # GRN columns that are bookkeeping rather than data a user would want in a sheet
    skip_grn = {"id", "lot_id", "file_path", "file_hash", "is_deleted", "deleted_at", "source", "lot_number"}
    wanted_grn = {
        "grn_number": "grn number", "invoice_number": "invoice number", "invoice_date": "date of invoice",
        "sender_name": "vendor name", "quantity": "invoice quantity", "amount": "invoice price",
        "purchase_date": "purchase date", "grn_date": "grn date", "po_number": "po number",
        "vehicle_number": "vehicle number", "e_way_bill": "e-way bill", "notes": "grn notes",
        "validated": "validated", "received_qty": "received qty", "validation_ref": "validation ref",
        "validation_notes": "validation notes", "file_name": "invoice file",
        "created_by": "received by", "created_at": "grn created at",
    }
    for col in GRNImport.__table__.columns.keys():
        if col in skip_grn:
            continue
        assert col in wanted_grn, f"GRNImport.{col} neither exported nor deliberately skipped"
        assert wanted_grn[col] in headers, f"GRNImport.{col} missing from export"
    wanted_lot = {
        "lot_number": "lot number", "supplier_name": "lot supplier", "grn_system_number": "lot grn system number",
        "grn_number_new": None, "grn_date": "lot grn date", "invoice_date": "lot invoice date",
        "invoice_no": "lot invoice no", "invoice_value": "lot invoice value",
        "taxable_amount": "lot taxable amount", "sgst": "sgst", "cgst": "cgst", "igst": "igst",
        "vehicle_number": "lot vehicle number", "e_way_bill": "lot e-way bill", "po_number": "lot po number",
        "vendor_name": "lot vendor", "buying_price": "lot buying price", "qty": "lot qty",
        "purchase_date": "lot purchase date", "condition": "lot condition", "selling_price": "lot selling price",
        "notes": "lot notes", "created_by": "lot created by", "created_at": "lot created at",
        "is_restricted": "lot restricted",
    }
    skip_lot = {"id", "is_trashed", "trashed_at"}
    for col in Lot.__table__.columns.keys():
        if col in skip_lot:
            continue
        assert col in wanted_lot, f"Lot.{col} neither exported nor deliberately skipped"
        if wanted_lot[col]:
            assert wanted_lot[col] in headers, f"Lot.{col} missing from export"


# ── DB-backed: filter + export through the real routes ─────────────────────

def _run(src):
    r = subprocess.run([sys.executable, "-c", src], capture_output=True, text=True, cwd=ROOT, timeout=60)
    if r.returncode != 0:
        raise RuntimeError(f"subprocess failed:\n{r.stdout}\n{r.stderr}")
    return r.stdout.strip()


def _seed(rows, lot_number, lot_for_grn):
    """rows = [(grn_number, invoice_date, vendor, source)]; one Lot mapped onto lot_for_grn."""
    _run(f"""
import asyncio, sys
sys.path.insert(0, r"{ROOT}")
from datetime import datetime
from decimal import Decimal
from database import AsyncSessionLocal
from models.grn_import import GRNImport
from models.lot import Lot

async def main():
    async with AsyncSessionLocal() as db:
        for num, inv_date, vendor, source in {rows!r}:
            db.add(GRNImport(grn_number=num, source=source, created_by="itest-recv", sender_name=vendor,
                              invoice_date=inv_date, invoice_number="INV-" + num[-4:], quantity=7,
                              amount=Decimal("1234.50"), po_number="PO-" + num[-4:], notes="grn-note"))
        db.add(Lot(lot_number="{lot_number}", supplier_name="ITSupplier", vendor_name="ITLotVendor",
                    grn_system_number="{lot_for_grn}", invoice_no="LINV-1", invoice_value=Decimal("999.00"),
                    taxable_amount=Decimal("846.61"), sgst=Decimal("76.20"), cgst=Decimal("76.19"),
                    igst=Decimal("0"), buying_price=Decimal("800.00"), qty=7, e_way_bill="EWB-1",
                    vehicle_number="MH01AB1234", po_number="LPO-1", condition="Used", notes="lot-note",
                    created_by="itest-lot", purchase_date=datetime(2026, 1, 1)))
        await db.commit()

asyncio.run(main())
""")


def _grn_ids(numbers):
    out = _run(f"""
import asyncio, sys
sys.path.insert(0, r"{ROOT}")
from sqlalchemy import select
from database import AsyncSessionLocal
from models.grn_import import GRNImport

async def main():
    async with AsyncSessionLocal() as db:
        rows = (await db.execute(select(GRNImport).where(GRNImport.grn_number.in_({list(numbers)!r})))).scalars().all()
        for g in rows:
            print(g.grn_number, g.id)

asyncio.run(main())
""")
    return dict(line.split() for line in out.splitlines())


def _cleanup(numbers, lot_number):
    _run(f"""
import asyncio, sys
sys.path.insert(0, r"{ROOT}")
from sqlalchemy import select
from database import AsyncSessionLocal
from models.grn_import import GRNImport
from models.lot import Lot

async def main():
    async with AsyncSessionLocal() as db:
        lot = (await db.execute(select(Lot).where(Lot.lot_number == "{lot_number}"))).scalar_one_or_none()
        if lot:
            await db.delete(lot)
        for g in (await db.execute(select(GRNImport).where(GRNImport.grn_number.in_({list(numbers)!r})))).scalars().all():
            await db.delete(g)
        await db.commit()

asyncio.run(main())
""")


def _sheet(resp):
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(resp.content)).active
    rows = list(ws.iter_rows(values_only=True))
    headers = list(rows[0])
    return headers, [dict(zip(headers, r)) for r in rows[1:]]


def test_invoice_date_filter_and_export_selection(app_client, make_user):  # noqa: F811
    tag = str(uuid.uuid4().int)[:8]
    g_jan, g_feb, g_mar = f"9101{tag}", f"9102{tag}", f"9103{tag}"   # 12 digits each
    lot_number = f"ITGRNX{tag}"
    # Mixed free-text formats on purpose; the Feb GRN is the one with a Lot.
    _seed([(g_jan, "2026-01-10", "ITVendorJan", "invoice"),
           (g_feb, "15/02/2026", "ITVendorFeb", "invoice"),
           (g_mar, "2026-03-20", "ITVendorMar", "invoice")], lot_number, g_feb)
    try:
        username, password = make_user("admin")
        _login(app_client, username, password)
        csrf = app_client.cookies.get("csrf_token") or "dummy"
        ids = _grn_ids([g_jan, g_feb, g_mar])

        # 1) page: server-side Date of Invoice range
        page = app_client.get("/grn", params={"date_from": "2026-02-01", "date_to": "2026-02-28"})
        assert page.status_code == 200
        assert g_feb in page.text and g_jan not in page.text and g_mar not in page.text
        assert 'class="form-check-input grn-row-chk"' in page.text           # row checkbox
        assert 'id="grnSelectAll"' in page.text and "Export GRN" in page.text
        unfiltered = app_client.get("/grn")
        assert all(n in unfiltered.text for n in (g_jan, g_feb, g_mar))

        # 2) export with a selection -> only the selected GRN, with ALL fields
        r = app_client.post("/grn/export", data={"csrf_token": csrf, "source": "invoice",
                                                 "ids": ids[g_feb]})
        assert r.status_code == 200
        assert "spreadsheetml" in r.headers["content-type"]
        headers, rows = _sheet(r)
        assert [x["GRN Number"] for x in rows] == [g_feb]
        row = rows[0]
        assert row["Date of Invoice"] == "2026-02-15"            # normalised
        assert row["Vendor Name"] == "ITVendorFeb" and row["Invoice Quantity"] == 7
        assert row["Invoice Price (Amount)"] == 1234.5 and row["PO Number"] == f"PO-{g_feb[-4:]}"
        assert row["GRN Notes"] == "grn-note" and row["Received By (Created By)"] == "itest-recv"
        assert row["Lot Number"] == lot_number and row["Lot Supplier"] == "ITSupplier"
        assert row["Lot Vendor (PO)"] == "ITLotVendor" and row["Lot Invoice Value"] == 999.0
        assert row["SGST"] == 76.2 and row["CGST"] == 76.19 and row["Lot Qty"] == 7
        assert row["Lot E-Way Bill"] == "EWB-1" and row["Lot Condition"] == "Used"
        assert row["Lot Notes"] == "lot-note" and row["Lot Created By"] == "itest-lot"
        for needed in ("Invoice Number", "GRN Date", "Purchase Date", "Validated", "Lot Buying Price"):
            assert needed in headers

        # 3) two ticked rows -> exactly those two (a GRN without a Lot still gets a row)
        r = app_client.post("/grn/export", data={"csrf_token": csrf, "source": "invoice",
                                                 "ids": f"{ids[g_jan]},{ids[g_mar]}"})
        _, rows = _sheet(r)
        assert sorted(x["GRN Number"] for x in rows) == sorted([g_jan, g_mar])
        assert all(x["Lot Number"] is None for x in rows)

        # 4) nothing ticked -> everything matching the date range
        r = app_client.post("/grn/export", data={"csrf_token": csrf, "source": "invoice", "ids": "",
                                                 "date_from": "2026-01-01", "date_to": "2026-02-28"})
        nums = {x["GRN Number"] for x in _sheet(r)[1]}
        assert {g_jan, g_feb} <= nums and g_mar not in nums

        # 5) selection + range combine (selected but outside the range is dropped)
        r = app_client.post("/grn/export", data={"csrf_token": csrf, "source": "invoice",
                                                 "ids": f"{ids[g_jan]},{ids[g_feb]}",
                                                 "date_from": "2026-02-01", "date_to": "2026-02-28"})
        assert [x["GRN Number"] for x in _sheet(r)[1]] == [g_feb]

        # 6) CSV variant
        r = app_client.post("/grn/export", data={"csrf_token": csrf, "source": "invoice",
                                                 "ids": ids[g_feb], "fmt": "csv"})
        assert r.headers["content-type"].startswith("text/csv")
        assert g_feb in r.text and lot_number in r.text
    finally:
        _cleanup([g_jan, g_feb, g_mar], lot_number)


def test_post_iqc_page_filters_and_scopes_export_by_source(app_client, make_user):  # noqa: F811
    tag = str(uuid.uuid4().int)[:8]
    g_in, g_out, g_inv = f"9201{tag}", f"9202{tag}", f"9203{tag}"
    lot_number = f"ITGRNY{tag}"
    _seed([(g_in, "2026-04-10", "ITPostIn", "post_iqc"),
           (g_out, "2026-06-10", "ITPostOut", "post_iqc"),
           (g_inv, "2026-04-11", "ITInvoiceSrc", "invoice")], lot_number, g_in)
    try:
        username, password = make_user("admin")
        _login(app_client, username, password)
        csrf = app_client.cookies.get("csrf_token") or "dummy"

        page = app_client.get("/grn/post-iqc", params={"date_from": "2026-04-01", "date_to": "2026-04-30"})
        assert page.status_code == 200
        assert g_in in page.text and g_out not in page.text and g_inv not in page.text

        r = app_client.post("/grn/export", data={"csrf_token": csrf, "source": "post_iqc", "ids": "",
                                                 "date_from": "2026-04-01", "date_to": "2026-04-30"})
        rows = _sheet(r)[1]
        nums = {x["GRN Number"] for x in rows}
        assert g_in in nums and g_out not in nums and g_inv not in nums
        assert {x["GRN Source"] for x in rows} == {"GRN post IQC"}
        assert [x["Lot Number"] for x in rows if x["GRN Number"] == g_in] == [lot_number]
    finally:
        _cleanup([g_in, g_out, g_inv], lot_number)


def test_manage_lots_grn_tab_has_checkboxes_date_filter_and_export_route(app_client, make_user):  # noqa: F811
    tag = str(uuid.uuid4().int)[:8]
    g_a, g_b = f"9301{tag}", f"9302{tag}"
    lot_number = f"ITGRNZ{tag}"
    from models.master import EXTERNAL_PARTNER_TEST_GRN_SOURCE as SRC
    _seed([(g_a, "2026-05-02", "ITTestA", SRC), (g_b, "2026-07-02", "ITTestB", SRC)], lot_number, g_a)
    try:
        username, password = make_user("admin")
        _login(app_client, username, password)
        csrf = app_client.cookies.get("csrf_token") or "dummy"

        page = app_client.get("/trade-partner/manage-lots",
                              params={"date_from": "2026-05-01", "date_to": "2026-05-31"})
        assert page.status_code == 200
        assert g_a in page.text and g_b not in page.text
        assert 'id="grnSelectAll"' in page.text and 'id="exportGrnBtn"' in page.text
        assert "Date of Invoice" in page.text

        r = app_client.post("/trade-partner/manage-lots/export-grn",
                            data={"csrf_token": csrf, "ids": "",
                                  "date_from": "2026-05-01", "date_to": "2026-05-31"})
        assert r.status_code == 200
        rows = _sheet(r)[1]
        assert [x["GRN Number"] for x in rows] == [g_a]
        assert rows[0]["Lot Number"] == lot_number
    finally:
        _cleanup([g_a, g_b], lot_number)
